"""Sinh file Office mẫu good/bad để kiểm chứng convention MarkItDown."""
import base64
import io
from pathlib import Path

SAMPLES_DIR = Path(__file__).parent

# PNG 1x1 hợp lệ — đủ để python-docx/python-pptx nhúng làm ảnh
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def _repeat_header(table):
    """Đánh dấu dòng đầu 'Repeat as header row' -> mammoth map thành <th>."""
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    el = parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>')
    trPr.append(el)


def _add_hyperlink(paragraph, url, text):
    """Hyperlink thật -> [text](url) trong Markdown (DOCX-12)."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    xml = (
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}">'
        f'<w:r><w:rPr/><w:t xml:space="preserve">{text}</w:t></w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(parse_xml(xml))


def _add_altcontent_textbox(paragraph, text):
    """Text box kiểu Word thật (mc:AlternateContent: Choice DrawingML + Fallback VML).
    mammoth đọc nhánh Fallback -> chữ SỐNG SÓT nhưng bị dồn ra cuối đoạn (rối thứ tự)."""
    xml = (
        f'<w:r {nsdecls("w", "wp", "a")} '
        f'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        f'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        f'xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<mc:AlternateContent><mc:Choice Requires="wps">'
        f'<w:drawing><wp:inline><wp:extent cx="2000000" cy="500000"/>'
        f'<wp:docPr id="11" name="TextBox"/>'
        f'<a:graphic><a:graphicData '
        f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r>'
        f'<w:t xml:space="preserve">{text}</w:t>'
        f'</w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>'
        f'</a:graphicData></a:graphic></wp:inline></w:drawing>'
        f'</mc:Choice><mc:Fallback>'
        f'<w:pict><v:shape style="width:220pt;height:36pt"><v:textbox>'
        f'<w:txbxContent><w:p><w:r>'
        f'<w:t xml:space="preserve">{text}</w:t>'
        f'</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>'
        f'</mc:Fallback></mc:AlternateContent></w:r>'
    )
    paragraph._p.append(parse_xml(xml))


def _add_drawingml_graphic(paragraph, text):
    """Đồ hoạ DrawingML KHÔNG có VML fallback (đại diện SmartArt: chữ ngoài tầm đọc
    của mammoth) -> chữ MẤT hoàn toàn."""
    xml = (
        f'<w:r {nsdecls("w", "wp", "a")} '
        f'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="2000000" cy="500000"/>'
        f'<wp:docPr id="12" name="Graphic"/>'
        f'<a:graphic><a:graphicData '
        f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r>'
        f'<w:t xml:space="preserve">{text}</w:t>'
        f'</w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>'
        f'</a:graphicData></a:graphic>'
        f'</wp:inline></w:drawing></w:r>'
    )
    paragraph._p.append(parse_xml(xml))


def _add_omml_pythagoras(paragraph):
    """Chèn công thức OMML a^2 + b^2 -> MarkItDown chuyển thành LaTeX (DOCX-05)."""
    xml = (
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:sSup><m:e><m:r><m:t>a</m:t></m:r></m:e>'
        '<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
        '<m:r><m:t xml:space="preserve"> + </m:t></m:r>'
        '<m:sSup><m:e><m:r><m:t>b</m:t></m:r></m:e>'
        '<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
        '</m:oMath>'
    )
    paragraph._p.append(parse_xml(xml))


def _add_tracked_insertion(paragraph, text):
    """Insertion đang treo (chưa Accept) -> tracked change chưa resolve (DOCX-06)."""
    xml = (
        f'<w:ins {nsdecls("w")} w:id="1" w:author="reviewer" '
        f'w:date="2026-07-15T00:00:00Z"><w:r>'
        f'<w:t xml:space="preserve">{text}</w:t></w:r></w:ins>'
    )
    paragraph._p.append(parse_xml(xml))


def gen_docx() -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    # ---------- GOOD ----------
    doc = Document()
    doc.add_heading("Báo cáo doanh thu 2026 [DOCX-01]", level=1)      # DOCX-01
    doc.add_heading("Tình hình quý 1 [DOCX-01]", level=2)
    p9 = doc.add_paragraph(
        "Mã dự án DT-2026, phiên bản 1.0 — định danh đặt ở đầu thân bài, "
        "không nhét vào header/footer. [DOCX-09]"                     # DOCX-09
    )
    doc.add_paragraph(
        "Doanh thu quý 1 đạt 600 triệu (số liệu gõ thành văn bản, không chỉ "
        "nằm trong ảnh). [DOCX-04]"                                   # DOCX-04
    )
    p5 = doc.add_paragraph("Ràng buộc kiểm tra [DOCX-05]: ")          # DOCX-05
    _add_omml_pythagoras(p5)

    table = doc.add_table(rows=3, cols=2)                             # DOCX-02
    data = [("Tháng [DOCX-02]", "Doanh thu"), ("01", "100"), ("02", "200")]
    for row, (c1, c2) in zip(table.rows, data):
        row.cells[0].text = c1
        row.cells[1].text = c2
    _repeat_header(table)

    doc.add_paragraph("Bước khảo sát khách hàng [DOCX-11]", style="List Bullet")
    doc.add_paragraph("Bước chốt tính năng [DOCX-11]", style="List Bullet")  # DOCX-11

    p12 = doc.add_paragraph("Nguồn tham chiếu [DOCX-12]: ")           # DOCX-12
    _add_hyperlink(p12, "https://example.com/bao-cao", "báo cáo đầy đủ")

    p10 = doc.add_paragraph(
        "Lưu ý: số liệu quý 1 là tạm tính, chốt lại cuối tháng 4. "
        "(ghi chú giữ lại đặt trong thân bài, không dùng comment) [DOCX-10]"  # DOCX-10
    )
    # DOCX-07: thay text box bằng bảng 1 ô để nhấn mạnh
    box_tbl = doc.add_table(rows=1, cols=1)
    box_tbl.style = "Table Grid"
    box_tbl.rows[0].cells[0].text = (
        "KHUNG NHẤN MẠNH: dùng bảng 1 ô thay cho text box. [DOCX-07]"
    )
    # DOCX-08: thay SmartArt bằng bullet list
    doc.add_paragraph("Quy trình thay cho SmartArt [DOCX-08]:")
    doc.add_paragraph("B1 khảo sát", style="List Number")
    doc.add_paragraph("B2 phân tích", style="List Number")
    # DOCX-03: ảnh có alt text
    pic = doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))
    pic._inline.docPr.set(
        "descr",
        "Biểu đồ cột doanh thu theo tháng: tháng 1 đạt 100, tháng 2 đạt 200 [DOCX-03]",
    )
    pic._inline.docPr.set("title", "Biểu đồ doanh thu")
    # DOCX-06: file good không còn tracked change (chỉ ghi chú)
    doc.add_paragraph(
        "Đã Accept toàn bộ thay đổi, tắt Track Changes trước khi nộp. [DOCX-06]"
    )
    doc.save(SAMPLES_DIR / "docx-good.docx")

    # ---------- BAD ----------
    doc = Document()
    p = doc.add_paragraph()                                          # DOCX-01 (vi phạm)
    run = p.add_run("Báo cáo doanh thu 2026 [vi phạm DOCX-01: bold+cỡ chữ]")
    run.bold = True
    run.font.size = Pt(16)

    # DOCX-09 (vi phạm): nội dung đặt trong header
    doc.sections[0].header.paragraphs[0].text = (
        "Mã dự án DT-2026 [vi phạm DOCX-09: đặt trong header, sẽ biến mất]"
    )

    # DOCX-02 (vi phạm): bảng merge header, không đánh dấu header row
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].merge(table.rows[0].cells[1])
    table.rows[0].cells[0].text = "Doanh thu theo tháng [vi phạm DOCX-02: merge]"
    table.rows[1].cells[0].text = "01"
    table.rows[1].cells[1].text = "100"
    table.rows[2].cells[0].text = "02"
    table.rows[2].cells[1].text = "200"

    # DOCX-05 (vi phạm): "công thức" là ảnh, không có bản OMML
    doc.add_paragraph(
        "Ràng buộc a^2+b^2 chỉ tồn tại trong ảnh dưới đây [vi phạm DOCX-05]:"
    )
    doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))

    # DOCX-07 (vi phạm): text box -> chữ sống sót nhưng bị dồn ra cuối đoạn
    p7 = doc.add_paragraph("Doanh thu quý 1: ")
    _add_altcontent_textbox(p7, "ĐẠT 600 TRIỆU")
    p7.add_run(" — đã kiểm toán. [vi phạm DOCX-07]")

    # DOCX-08 (vi phạm): SmartArt/đồ hoạ DrawingML -> chữ MẤT hoàn toàn
    doc.add_paragraph(
        "Sơ đồ quy trình dưới đây là SmartArt/đồ hoạ, chữ bên trong sẽ MẤT "
        "[vi phạm DOCX-08]:"
    )
    p8 = doc.add_paragraph()
    _add_drawingml_graphic(p8, "KHẢO SÁT → PHÂN TÍCH → CHỐT")

    # DOCX-06 (vi phạm): tracked change chưa resolve
    p6 = doc.add_paragraph("Câu có sửa đổi đang treo [vi phạm DOCX-06]: ")
    _add_tracked_insertion(p6, "PHẦN CHÈN ĐANG TREO")

    # DOCX-11 (vi phạm): gõ tay dấu gạch đầu dòng
    doc.add_paragraph("- Bước khảo sát (gõ tay) [vi phạm DOCX-11]")
    doc.add_paragraph("- Bước chốt tính năng (gõ tay)")

    # DOCX-03/04 (vi phạm): ảnh không alt text, số liệu chỉ nằm trong ảnh
    doc.add_paragraph(
        "Bảng số liệu bên dưới chỉ là ảnh chụp, không có bản chữ [vi phạm DOCX-04]:"
    )
    doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))
    doc.save(SAMPLES_DIR / "docx-bad.docx")


def gen_xlsx() -> None:
    from openpyxl import Workbook

    # good: 1 sheet = 1 bảng từ A1, dòng 1 là header, không merge,
    # dòng Tổng ghi GIÁ TRỊ chứ không ghi công thức
    wb = Workbook()
    ws = wb.active
    ws.title = "DoanhThu2026"
    ws.append(["Tháng", "Doanh thu", "Trạng thái"])
    ws.append(["01", 100, "Đạt"])
    ws.append(["02", 200, "Đạt"])
    ws.append(["03", 300, "Vượt"])
    # "Không áp dụng" thay vì "N/A": pandas coi literal "N/A" là giá trị khuyết -> NaN
    ws.append(["Tổng", 600, "Không áp dụng"])
    wb.save(SAMPLES_DIR / "xlsx-good.xlsx")

    # bad: tiêu đề trang trí + merge, bảng lệch khỏi A1, ô bỏ trống,
    # công thức chưa từng được Excel tính (không cached value), sheet ẩn
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "BÁO CÁO DOANH THU NĂM 2026"
    ws.merge_cells("A1:C1")
    ws["A3"] = "Tháng"
    ws["B3"] = "Doanh thu"
    ws["C3"] = "Ghi chú"
    ws["A4"] = "01"
    ws["B4"] = 110
    ws["C4"] = "N/A"         # literal "N/A" -> pandas coi là giá trị khuyết -> NaN
    ws["A5"] = "02"          # B5 cố tình bỏ trống
    ws["A6"] = "Gấp ba"
    ws["B6"] = "=B4*3"       # openpyxl không tính -> không có cached value
    hidden = wb.create_sheet("NhapLieuTam")
    hidden["A1"] = "dữ liệu nháp không nên lộ ra"
    hidden.sheet_state = "hidden"
    wb.save(SAMPLES_DIR / "xlsx-bad.xlsx")


def gen_pptx() -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    # good: layout có Title placeholder, bullet trong body, ảnh có alt text,
    # speaker notes, chart cột cơ bản
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Kế hoạch quý 3"
    body = slide.placeholders[1].text_frame
    body.text = "Bước 1: khảo sát khách hàng"
    body.add_paragraph().text = "Bước 2: chốt tính năng"
    pic = slide.shapes.add_picture(
        io.BytesIO(PNG_1PX), Inches(6), Inches(1), width=Inches(1)
    )
    pic._element._nvXxPr.cNvPr.set(
        "descr", "Sơ đồ timeline quý 3: tháng 7 khảo sát, tháng 8 chốt tính năng"
    )
    slide.notes_slide.notes_text_frame.text = (
        "Chi tiết: khảo sát 50 khách hàng nhóm A trong tháng 7."
    )
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide2.shapes.title.text = "Doanh thu theo tháng"
    chart_data = CategoryChartData()
    chart_data.categories = ["Tháng 7", "Tháng 8"]
    chart_data.add_series("Doanh thu", (100, 200))
    slide2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(2), Inches(6), Inches(4), chart_data,
    )
    prs.save(SAMPLES_DIR / "pptx-good.pptx")

    # bad: slide Blank, "tiêu đề" là textbox nằm ĐÁY slide, các bước đặt
    # đảo vị trí (Bước 2 ở trên, Bước 1 ở dưới), ảnh không alt text
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    box2 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box2.text_frame.text = "Bước 2: chốt tính năng"
    box1 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    box1.text_frame.text = "Bước 1: khảo sát khách hàng"
    title_box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    title_box.text_frame.text = "Kế hoạch quý 3"
    pic_bad = slide.shapes.add_picture(
        io.BytesIO(PNG_1PX), Inches(6), Inches(5), width=Inches(1)
    )
    # python-pptx tự viết descr="image.png" khi add_picture từ stream — đó là
    # artifact của python-pptx, KHÔNG phải hành vi MarkItDown. Xóa để mẫu phản
    # ánh đúng trường hợp "ảnh không có alt text" (descr rỗng) mà convention mô tả.
    pic_bad._element._nvXxPr.cNvPr.set("descr", "")
    prs.save(SAMPLES_DIR / "pptx-bad.pptx")


def main() -> None:
    (SAMPLES_DIR / "output").mkdir(exist_ok=True)
    gen_docx()
    print("generated: docx")
    gen_xlsx()
    print("generated: xlsx")
    gen_pptx()
    print("generated: pptx")


if __name__ == "__main__":
    main()
